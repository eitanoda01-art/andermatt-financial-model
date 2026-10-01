import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Protection, PatternFill, Font
from openpyxl.comments import Comment
import shutil

SRC = "/root/.claude/uploads/ca8257f0-1581-55bf-88a4-34c1d47ea869/fb5c87e5-Aura_Andermatt_10Year_Model.xlsx"
DST = "/home/user/andermatt-financial-model/Aura_Andermatt_10Year_Model.xlsx"

shutil.copy2(SRC, DST)
wb = openpyxl.load_workbook(DST)

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# REVENUE SHEET CHANGES
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
rev = wb["Revenue"]

# ── 1. Rename "GLOBAL REVENUE VARIABLES" → "MODEL ASSUMPTIONS" ──
rev["A3"].value = "MODEL ASSUMPTIONS"
ASSUMPTIONS_FILL = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")
ASSUMPTIONS_FONT = Font(bold=True, size=11, color="1F4E79")
rev["A3"].fill = ASSUMPTIONS_FILL
rev["A3"].font = ASSUMPTIONS_FONT

YELLOW = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
SECTION_FILL = PatternFill(start_color="FFD9E2F3", end_color="FFD9E2F3", fill_type="solid")
SECTION_FONT = Font(bold=True, size=11)
COMMENT_FONT = Font(size=9, color="666666")
LABEL_FONT = Font(size=11)

COL_LETTERS = ["", "A","B","C","D","E","F","G","H","I","J","K"]
YEARS = list(range(2027, 2037))

def clear_rows(ws, start_row, end_row):
    for r in range(start_row, end_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell.fill = PatternFill(fill_type=None)
            cell.font = Font()
            cell.comment = None

def write_year_headers(ws, row):
    for i, year in enumerate(YEARS):
        ws.cell(row=row, column=i+2).value = year
        ws.cell(row=row, column=i+2).font = Font(bold=True)

def write_section_header(ws, row, title):
    ws.cell(row=row, column=1).value = title
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.cell(row=row, column=1).font = SECTION_FONT

def write_comment_row(ws, row, text):
    ws.cell(row=row, column=1).value = text
    ws.cell(row=row, column=1).font = COMMENT_FONT

def write_yellow_cell(ws, row, col, value):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.fill = YELLOW

def write_fb_block(ws, base_row, label_prefix, guest_count, per_person, share_rate):
    """Write a 5-row F&B sub-block: guest count, spend, share, est total, revenue."""
    r = base_row
    ws.cell(row=r, column=1).value = f"  {label_prefix} — Guest Count"
    ws.cell(row=r, column=1).font = LABEL_FONT
    write_yellow_cell(ws, r, 2, guest_count)

    r += 1
    ws.cell(row=r, column=1).value = f"  {label_prefix} — Per-person Spend"
    ws.cell(row=r, column=1).font = LABEL_FONT
    write_yellow_cell(ws, r, 2, per_person)

    r += 1
    ws.cell(row=r, column=1).value = f"  {label_prefix} — Share Rate"
    ws.cell(row=r, column=1).font = LABEL_FONT
    write_yellow_cell(ws, r, 2, share_rate)
    share_comment = Comment(
        "Share rate may be based on profit margin, not gross revenue. "
        "Adjust when contract terms are confirmed.",
        "Model Note"
    )
    share_comment.width = 300
    share_comment.height = 80
    ws.cell(row=r, column=2).comment = share_comment

    r += 1
    guest_row = base_row
    spend_row = base_row + 1
    ws.cell(row=r, column=1).value = f"  {label_prefix} — Est. Total Sales (reference)"
    ws.cell(row=r, column=1).font = LABEL_FONT
    ws.cell(row=r, column=2).value = f"=B{guest_row}*B{spend_row}"
    for i in range(1, 10):
        cl = COL_LETTERS[i+2]
        prev = COL_LETTERS[i+1]
        ws.cell(row=r, column=i+2).value = f"={prev}{r}*(1+$B$9)"

    r += 1
    est_row = r - 1
    share_row = base_row + 2
    ws.cell(row=r, column=1).value = f"  {label_prefix} — Aura Revenue"
    ws.cell(row=r, column=1).font = LABEL_FONT
    for i in range(10):
        cl = COL_LETTERS[i+2]
        ws.cell(row=r, column=i+2).value = f"={cl}{est_row}*$B{share_row}"

    return r  # last row used

# ── 2. Rewrite Bar/F&B section (rows 64+) ──
# Clear from row 64 to end of sheet
clear_rows(rev, 64, rev.max_row)

# Extend sheet if needed (original max was 101, new max is 114)
# Just writing to new rows will extend it automatically.

# Section header
ROW = 64
write_section_header(rev, ROW, "SECTION 4: BAR / F&B REVENUE")
ROW += 1
write_comment_row(rev, ROW, "  Share rates are currently set to 0%. Aura does not take a share of bar/food revenue. If a share agreement is negotiated, change the share rate.")
ROW += 2  # skip blank row
write_year_headers(rev, ROW)  # row 67
ROW += 1

# Friday Night: Coffee + Drinks (rows 68-72)
FB_FRIDAY_BASE = ROW  # 68
last = write_fb_block(rev, ROW, "Friday Night: Coffee + Drinks", 150, 50, 0)
FRIDAY_REV_ROW = last  # 72
ROW = last + 2  # skip blank

# Saturday Day: Drinks (rows 74-78)
FB_SATDAY_DRINKS_BASE = ROW  # 74
last = write_fb_block(rev, ROW, "Saturday Day: Drinks", 280, 20, 0)
SATDAY_DRINKS_REV_ROW = last  # 78
ROW = last + 2

# Saturday Day: Food (rows 80-84)
FB_SATDAY_FOOD_BASE = ROW  # 80
last = write_fb_block(rev, ROW, "Saturday Day: Food", 280, 50, 0)
SATDAY_FOOD_REV_ROW = last  # 84
ROW = last + 2

# Saturday Night: Drinks (rows 86-90)
FB_SATNIGHT_BASE = ROW  # 86
last = write_fb_block(rev, ROW, "Saturday Night: Drinks", 150, 50, 0)
SATNIGHT_REV_ROW = last  # 90
ROW = last + 2

# Sunday Coffee (rows 92-96)
FB_SUNDAY_BASE = ROW  # 92
last = write_fb_block(rev, ROW, "Sunday Coffee", 80, 8, 0)
SUNDAY_REV_ROW = last  # 96
ROW = last + 2

# Bar/F&B Total (row 98)
FB_TOTAL_ROW = ROW  # 98
rev.cell(row=ROW, column=1).value = "Bar/F&B Total"
rev.cell(row=ROW, column=1).font = Font(bold=True)
for i in range(10):
    cl = COL_LETTERS[i+2]
    rev.cell(row=ROW, column=i+2).value = (
        f"={cl}{FRIDAY_REV_ROW}+{cl}{SATDAY_DRINKS_REV_ROW}+"
        f"{cl}{SATDAY_FOOD_REV_ROW}+{cl}{SATNIGHT_REV_ROW}+{cl}{SUNDAY_REV_ROW}"
    )
ROW += 2

# ── Dinner Revenue ──
DINNER_HEADER_ROW = ROW  # 100
write_section_header(rev, ROW, "SECTION 5: DINNER REVENUE")
ROW += 1
write_year_headers(rev, ROW)  # 101
ROW += 1
DINNER_ROW = ROW  # 102
rev.cell(row=ROW, column=1).value = "Dinner Revenue (manual input)"
rev.cell(row=ROW, column=1).font = LABEL_FONT
for i in range(10):
    write_yellow_cell(rev, ROW, i+2, 0)
ROW += 1
write_comment_row(rev, ROW, "  BP: 'Dinner is invite-only, sponsor-funded, never sold.' If a dinner-specific sponsor appears, it goes under Partnerships.")
ROW += 2

# ── Other Revenue ──
OTHER_HEADER_ROW = ROW  # 105
write_section_header(rev, ROW, "SECTION 6: OTHER REVENUE")
ROW += 1
write_year_headers(rev, ROW)  # 106
ROW += 1
OTHER_ROW = ROW  # 107
rev.cell(row=ROW, column=1).value = "Other Revenue"
rev.cell(row=ROW, column=1).font = LABEL_FONT
write_yellow_cell(rev, ROW, 2, 2000)
for i in range(1, 10):
    cl = COL_LETTERS[i+2]
    prev = COL_LETTERS[i+1]
    rev.cell(row=ROW, column=i+2).value = f"={prev}{OTHER_ROW}*(1+'P&L'!$B$4)"
ROW += 1
write_comment_row(rev, ROW, "  Other Revenue grows with General Inflation Rate (P&L sheet). Contents: add-ons, upselling, merchandise.")
ROW += 2

# ── Revenue Totals ──
TOTALS_HEADER_ROW = ROW  # 110
rev.cell(row=ROW, column=1).value = "REVENUE TOTALS"
rev.cell(row=ROW, column=1).fill = SECTION_FILL
rev.cell(row=ROW, column=1).font = SECTION_FONT
ROW += 1
write_year_headers(rev, ROW)  # 111
ROW += 1

WINTER_ROW = ROW  # 112
rev.cell(row=ROW, column=1).value = "TOTAL WINTER REVENUE"
rev.cell(row=ROW, column=1).font = Font(bold=True)
for i in range(10):
    cl = COL_LETTERS[i+2]
    rev.cell(row=ROW, column=i+2).value = (
        f"={cl}40+{cl}56+{cl}61+{cl}{FB_TOTAL_ROW}+{cl}{DINNER_ROW}+{cl}{OTHER_ROW}"
    )
ROW += 1

SUMMER_ROW = ROW  # 113
rev.cell(row=ROW, column=1).value = "Summer Revenue"
rev.cell(row=ROW, column=1).font = Font(bold=True)
for i, year in enumerate(YEARS):
    cl = COL_LETTERS[i+2]
    rev.cell(row=ROW, column=i+2).value = f"=IF({year}>=$B$13,{cl}{WINTER_ROW}*$B$11,0)"
ROW += 1

ANNUAL_ROW = ROW  # 114
rev.cell(row=ROW, column=1).value = "TOTAL ANNUAL REVENUE"
rev.cell(row=ROW, column=1).font = Font(bold=True)
for i in range(10):
    cl = COL_LETTERS[i+2]
    rev.cell(row=ROW, column=i+2).value = f"={cl}{WINTER_ROW}+{cl}{SUMMER_ROW}"

# ── 3. Update P&L cross-references ──
pl = wb["P&L"]

# P&L R13: Bar/F&B Total (was Revenue!85 → now FB_TOTAL_ROW)
for i in range(10):
    cl = COL_LETTERS[i+2]
    pl.cell(row=13, column=i+2).value = f"=Revenue!{cl}{FB_TOTAL_ROW}"

# P&L R14: Dinner Revenue (was Revenue!89 → now DINNER_ROW)
for i in range(10):
    cl = COL_LETTERS[i+2]
    pl.cell(row=14, column=i+2).value = f"=Revenue!{cl}{DINNER_ROW}"

# P&L R15: Other Revenue (was Revenue!94 → now OTHER_ROW)
for i in range(10):
    cl = COL_LETTERS[i+2]
    pl.cell(row=15, column=i+2).value = f"=Revenue!{cl}{OTHER_ROW}"

# P&L R17: TOTAL ANNUAL REVENUE (was Revenue!101 → now ANNUAL_ROW)
for i in range(10):
    cl = COL_LETTERS[i+2]
    pl.cell(row=17, column=i+2).value = f"=Revenue!{cl}{ANNUAL_ROW}"


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# COST SHEET CHANGES
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
cost = wb["Cost"]

# ── 1. Rename "GLOBAL COST VARIABLES" → "MODEL ASSUMPTIONS" ──
cost["A3"].value = "MODEL ASSUMPTIONS"
cost["A3"].fill = ASSUMPTIONS_FILL
cost["A3"].font = ASSUMPTIONS_FONT

# ── 2. Add comment to F&B Credit cell explaining the 9,000 CHF ──
fb_credit_comment = Comment(
    "BP shows 'Bar/F&B (net) = 9,000 CHF' as revenue, but per Rox confirmation, "
    "Aura receives no share of bar/food sales. This 9,000 is a deposit-style credit: "
    "guests are expected to spend more than 9,000 at the venue, so the excess offsets "
    "venue costs. It is correctly recorded here as a cost reduction, not as revenue.",
    "Model Note"
)
fb_credit_comment.width = 350
fb_credit_comment.height = 120
cost.cell(row=17, column=1).comment = fb_credit_comment

# ── 3. Add TOTAL COST row after Founder Pay (row 55) ──
COST_TOTAL_ROW = 57
cost.cell(row=COST_TOTAL_ROW, column=1).value = "TOTAL COST"
cost.cell(row=COST_TOTAL_ROW, column=1).fill = SECTION_FILL
cost.cell(row=COST_TOTAL_ROW, column=1).font = SECTION_FONT
for i in range(10):
    cl = COL_LETTERS[i+2]
    cost.cell(row=COST_TOTAL_ROW, column=i+2).value = f"={cl}49+{cl}53+{cl}55"


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# DASHBOARD CHANGES (same as before)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
ws = wb["Dashboard"]

# ── Rename "Net Result" → "Profit After Tax" ──
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
    for cell in row:
        if isinstance(cell.value, str) and "Net Result" in cell.value:
            cell.value = cell.value.replace("Net Result", "Profit After Tax")

pl["A29"].value = "PROFIT AFTER TAX"

# ── Remove "10-Year Total Revenue" row from KEY METRICS ──
for c in range(1, 5):
    ws.cell(row=6, column=c).value = None

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# REWRITE SCENARIO ADJUSTMENTS (rows 55-72)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
clear_rows(ws, 55, 76)

ws.cell(row=55, column=1).value = "SCENARIO ADJUSTMENTS (change values below to adjust scenarios)"
ws.cell(row=55, column=1).fill = ASSUMPTIONS_FILL
ws.cell(row=55, column=1).font = ASSUMPTIONS_FONT

for c, lbl in [(2, "Downside"), (3, "Base"), (4, "Upside")]:
    ws.cell(row=56, column=c).value = lbl
    ws.cell(row=56, column=c).font = Font(bold=True)

ws.cell(row=58, column=1).value = "■ REVENUE FACTORS"
ws.cell(row=58, column=1).font = Font(bold=True)
ws.cell(row=59, column=1).value = "  These multiply each year's revenue line directly."
ws.cell(row=59, column=1).font = COMMENT_FONT

ADJ_ROWS = {
    60: ("Ticket Occupancy", 0.85, 1, 1),
    61: ("Ticket Price Factor", 0.90, 1, 1.10),
    62: ("Partnership Factor", 0.60, 1, 1.40),
    63: ("ASA Factor", 0.75, 1, 1.25),
}
for r, (label, ds, base, us) in ADJ_ROWS.items():
    ws.cell(row=r, column=1).value = label
    ws.cell(row=r, column=1).font = LABEL_FONT
    for c, val in [(2, ds), (3, base), (4, us)]:
        ws.cell(row=r, column=c).value = val
        ws.cell(row=r, column=c).fill = YELLOW

ws.cell(row=65, column=1).value = "■ COST FACTORS"
ws.cell(row=65, column=1).font = Font(bold=True)

ws.cell(row=66, column=1).value = "Direct Cost Factor"
ws.cell(row=66, column=1).font = LABEL_FONT
for c, val in [(2, 0.90), (3, 1), (4, 1)]:
    ws.cell(row=66, column=c).value = val
    ws.cell(row=66, column=c).fill = YELLOW

ws.cell(row=67, column=1).value = "Founder Pay Factor"
ws.cell(row=67, column=1).font = LABEL_FONT
for c, val in [(2, 0.50), (3, 1), (4, 1)]:
    ws.cell(row=67, column=c).value = val
    ws.cell(row=67, column=c).fill = YELLOW

ws.cell(row=68, column=1).value = "Downside Overhead (CHF)"
ws.cell(row=68, column=1).font = LABEL_FONT
ws.cell(row=68, column=2).value = 8000
ws.cell(row=68, column=2).fill = YELLOW

ws.cell(row=70, column=1).value = "■ STRUCTURAL FACTORS"
ws.cell(row=70, column=1).font = Font(bold=True)

ws.cell(row=71, column=1).value = "Summer Start Year"
ws.cell(row=71, column=1).font = LABEL_FONT
for c, val in [(2, 2029), (3, 2027), (4, 2027)]:
    ws.cell(row=71, column=c).value = val
    ws.cell(row=71, column=c).fill = YELLOW

ws.cell(row=73, column=1).value = "SCENARIO DEFINITIONS:"
ws.cell(row=73, column=1).font = Font(bold=True, size=9, color="666666")
ws.cell(row=74, column=1).value = "  Downside: Tickets -15%, prices -10%, partnerships -40%, ASA -25%, direct costs ×0.9, overhead fixed 8K, founder pay halved (BP), summer from 2029."
ws.cell(row=74, column=1).font = COMMENT_FONT
ws.cell(row=75, column=1).value = "  Base: All factors at 1.0 — as entered in Revenue/Cost sheets. Summer from 2027."
ws.cell(row=75, column=1).font = COMMENT_FONT
ws.cell(row=76, column=1).value = "  Upside: Prices +10%, partnerships +40%, ASA +25%. Summer from 2027."
ws.cell(row=76, column=1).font = COMMENT_FONT

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# REWRITE ALL SCENARIO FORMULAS
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
SCENARIOS = [
    # (section_label, first_row, adj_col)  adj_col: B=Downside, C=Base, D=Upside
    ("BASE SCENARIO", 10, "C"),
    ("DOWNSIDE SCENARIO", 22, "B"),
    ("UPSIDE SCENARIO", 34, "D"),
]

for section_label, sec_start, adj_col in SCENARIOS:
    r_ticket = sec_start + 2   # 12, 24, 36
    r_partner = sec_start + 3  # 13, 25, 37
    r_other = sec_start + 4    # 14, 26, 38
    r_total_rev = sec_start + 5  # 15, 27, 39
    r_total_cost = sec_start + 6  # 16, 28, 40
    r_ebit = sec_start + 7     # 17, 29, 41
    r_tax = sec_start + 8      # 18, 30, 42
    r_pat = sec_start + 9      # 19, 31, 43
    r_cumcash = sec_start + 10  # 20, 32, 44

    # Labels (adj. suffix for non-base)
    suffix = " (adj.)" if adj_col != "C" else ""
    ws.cell(row=r_ticket, column=1).value = f"    Net Ticket Revenue{suffix}"
    ws.cell(row=r_partner, column=1).value = f"    Partnership Revenue{suffix}"
    ws.cell(row=r_other, column=1).value = f"    Other Revenue (ASA+F&B+Dinner+Other){suffix}"
    ws.cell(row=r_total_rev, column=1).value = "Total Revenue"
    ws.cell(row=r_total_cost, column=1).value = f"  Total Costs{suffix}"

    for i in range(10):
        cl = COL_LETTERS[i + 2]
        year = YEARS[i]

        # Ticket = P&L ticket × occupancy × price factor
        ws.cell(row=r_ticket, column=i+2).value = (
            f"='P&L'!{cl}10*${adj_col}$60*${adj_col}$61"
        )

        # Partnership = P&L partnership × partnership factor
        ws.cell(row=r_partner, column=i+2).value = (
            f"='P&L'!{cl}11*${adj_col}$62"
        )

        # Other = ASA×factor + F&B + Dinner + Other
        ws.cell(row=r_other, column=i+2).value = (
            f"='P&L'!{cl}12*${adj_col}$63+'P&L'!{cl}13+'P&L'!{cl}14+'P&L'!{cl}15"
        )

        # Total Revenue = (ticket+partner+other) × (1 + summer multiplier if year >= start)
        ws.cell(row=r_total_rev, column=i+2).value = (
            f"=({cl}{r_ticket}+{cl}{r_partner}+{cl}{r_other})"
            f"*(1+IF({year}>=${adj_col}$71,Revenue!$B$11,0))"
        )

        # Total Costs: Downside uses fixed overhead ($B$68), others use P&L overhead
        if adj_col == "B":  # Downside
            ws.cell(row=r_total_cost, column=i+2).value = (
                f"='P&L'!{cl}20*$B$66-$B$68+'P&L'!{cl}22*$B$67"
            )
        else:  # Base (C) / Upside (D)
            ws.cell(row=r_total_cost, column=i+2).value = (
                f"='P&L'!{cl}20*${adj_col}$66+'P&L'!{cl}21+'P&L'!{cl}22*${adj_col}$67"
            )

        # EBIT = Revenue + Costs (costs are negative)
        ws.cell(row=r_ebit, column=i+2).value = (
            f"={cl}{r_total_rev}+{cl}{r_total_cost}"
        )

        # Tax
        ws.cell(row=r_tax, column=i+2).value = (
            f"=IF({cl}{r_ebit}>0,-{cl}{r_ebit}*'P&L'!$B$5,0)"
        )

        # PAT
        ws.cell(row=r_pat, column=i+2).value = (
            f"={cl}{r_ebit}+{cl}{r_tax}"
        )

        # Cumulative Cash
        if i == 0:
            ws.cell(row=r_cumcash, column=i+2).value = (
                f"='P&L'!$B$6+{cl}{r_pat}"
            )
        else:
            prev_cl = COL_LETTERS[i + 1]
            ws.cell(row=r_cumcash, column=i+2).value = (
                f"={prev_cl}{r_cumcash}+{cl}{r_pat}"
            )

# ── Replace charts: remove all 3, add 2 new ones ──
ws._charts.clear()

cats = Reference(ws, min_col=2, max_col=11, min_row=11)

NAVY = "1F4E79"
RED = "C62828"
GREEN = "2E7D32"

def make_chart(ws, title, series_defs, cats):
    chart = LineChart()
    chart.title = title
    chart.style = 2
    chart.height = 13
    chart.width = 22
    chart.legend.position = 'b'

    chart.y_axis.numFmt = '#,##0'
    chart.y_axis.title = "CHF"
    chart.x_axis.title = None
    chart.y_axis.minorGridlines = None

    for row_num, label, color, dash, width in series_defs:
        ref = Reference(ws, min_col=2, max_col=11, min_row=row_num)
        chart.add_data(ref, from_rows=True)
        idx = len(chart.series) - 1
        s = chart.series[idx]
        s.title = openpyxl.chart.series.SeriesLabel(v=label)
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = width
        if dash:
            s.graphicalProperties.line.dashStyle = dash

        s.dLbls = DataLabelList()
        s.dLbls.showVal = True
        s.dLbls.numFmt = '#,##0'
        s.dLbls.showCatName = False
        s.dLbls.showSerName = False

    chart.set_categories(cats)
    return chart

chart1 = make_chart(ws, "Profit After Tax (CHF)", [
    (19, "Base", NAVY, None, 28000),
    (31, "Downside", RED, "dash", 18000),
    (43, "Upside", GREEN, "dash", 18000),
], cats)

chart2 = make_chart(ws, "Cumulative Cash (CHF)", [
    (20, "Base", NAVY, None, 28000),
    (32, "Downside", RED, "dash", 18000),
    (44, "Upside", GREEN, "dash", 18000),
], cats)

ws.add_chart(chart1, "M1")
ws.add_chart(chart2, "M17")


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# INSTRUCTIONS SHEET
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
if "Instructions" in wb.sheetnames:
    del wb["Instructions"]
ins = wb.create_sheet("Instructions", 0)

ins.column_dimensions["A"].width = 90
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
H2_FONT = Font(bold=True, size=12, color="1F4E79")
BODY_FONT = Font(size=11)
YELLOW_REF_FONT = Font(size=11, bold=True, color="997A00")

content = [
    ("AURA ANDERMATT — 10-Year Financial Model", TITLE_FONT),
    ("", None),
    ("HOW TO USE THIS MODEL", H2_FONT),
    ("", None),
    ("This workbook contains a 10-year financial simulation (2027–2036) for Aura Andermatt.", BODY_FONT),
    ("It consists of 4 working sheets: Dashboard, P&L, Revenue, and Cost.", BODY_FONT),
    ("", None),
    ("YELLOW CELLS = EDITABLE INPUTS", H2_FONT),
    ("", None),
    ("All sheets are protected. Only yellow-highlighted cells can be edited.", BODY_FONT),
    ("Yellow cells contain the key assumptions you can adjust:", BODY_FONT),
    ("  • Revenue sheet: ticket prices, guest counts, F&B spend, partnership values, growth rates", BODY_FONT),
    ("  • Cost sheet: individual cost items (venue, talent, security, etc.), overhead, founder pay", BODY_FONT),
    ("  • Dashboard: scenario adjustment factors (occupancy, price, partnership, ASA, costs, summer start)", BODY_FONT),
    ("All other cells contain formulas that update automatically when you change inputs.", BODY_FONT),
    ("", None),
    ("UNLOCKING PROTECTED CELLS", H2_FONT),
    ("", None),
    ("If you need to edit a locked (non-yellow) cell or modify the structure:", BODY_FONT),
    ("  1. Go to Review → Unprotect Sheet", BODY_FONT),
    ("  2. Enter password: 2027", BODY_FONT),
    ("  3. Make your changes", BODY_FONT),
    ("  4. Re-protect via Review → Protect Sheet (password: 2027) when done", BODY_FONT),
    ("", None),
    ("SHEET GUIDE", H2_FONT),
    ("", None),
    ("Dashboard", Font(bold=True, size=11)),
    ("  The summary view. Shows Base, Downside, and Upside scenarios side by side.", BODY_FONT),
    ("  Key Metrics (row 3–5): 10-year cumulative cash and 2027 PAT at a glance.", BODY_FONT),
    ("  Scenario tables (rows 10–44): Revenue, costs, EBIT, tax, PAT, cumulative cash per year.", BODY_FONT),
    ("  Scenario Adjustments (rows 55+): Change these yellow cells to adjust scenario assumptions.", BODY_FONT),
    ("  Two charts visualize PAT and cumulative cash across all three scenarios.", BODY_FONT),
    ("", None),
    ("P&L (Profit & Loss)", Font(bold=True, size=11)),
    ("  Consolidates Revenue and Cost sheets into a single annual P&L statement.", BODY_FONT),
    ("  All values are formula-driven from Revenue and Cost sheets — no manual input here.", BODY_FONT),
    ("  This represents the BASE case (no scenario adjustments applied).", BODY_FONT),
    ("", None),
    ("Revenue", Font(bold=True, size=11)),
    ("  Detailed revenue buildup: tickets (3 drops), partnerships, ASA, Bar/F&B, dinner, other.", BODY_FONT),
    ("  Model Assumptions (rows 3–13): growth rates, summer multiplier, net factor.", BODY_FONT),
    ("  Summer edition revenue = Winter revenue × Summer Multiplier (default 0.55).", BODY_FONT),
    ("", None),
    ("Cost", Font(bold=True, size=11)),
    ("  Three cost categories: A (fixed/contract), B (guest-linked), C (ambition-linked).", BODY_FONT),
    ("  Contingency % (row 6) is applied to direct costs. Summer costs = Winter × 0.5.", BODY_FONT),
    ("  Overhead and Founder Pay are separate from direct costs, entered manually per year.", BODY_FONT),
    ("", None),
    ("SCENARIO ADJUSTMENT FACTORS (Dashboard rows 55+)", H2_FONT),
    ("", None),
    ("  Factor                  Downside    Base    Upside", Font(size=11, name="Courier New")),
    ("  Ticket Occupancy          0.85      1.00      1.00", Font(size=11, name="Courier New")),
    ("  Ticket Price Factor       0.90      1.00      1.10", Font(size=11, name="Courier New")),
    ("  Partnership Factor        0.60      1.00      1.40", Font(size=11, name="Courier New")),
    ("  ASA Factor                0.75      1.00      1.25", Font(size=11, name="Courier New")),
    ("  Direct Cost Factor        0.90      1.00      1.00", Font(size=11, name="Courier New")),
    ("  Founder Pay Factor        0.50      1.00      1.00", Font(size=11, name="Courier New")),
    ("  Downside Overhead      8,000 CHF    (uses Cost sheet values)", Font(size=11, name="Courier New")),
    ("  Summer Start Year        2029      2027      2027", Font(size=11, name="Courier New")),
    ("", None),
    ("Base = P&L values as-is. Downside/Upside apply multipliers to the Base values.", BODY_FONT),
    ("To test custom scenarios, change the yellow factor cells on the Dashboard.", BODY_FONT),
    ("", None),
    ("NOTES", H2_FONT),
    ("", None),
    ("• F&B revenue share is currently set to 0% across all blocks (Aura receives no F&B revenue).", BODY_FONT),
    ("  If a share agreement is reached, update the Share Rate cells in the Revenue sheet.", BODY_FONT),
    ("• The F&B Credit (-9,000 CHF) on the Cost sheet is a deposit against guest F&B spending.", BODY_FONT),
    ("  See the cell comment on Cost!A17 for details.", BODY_FONT),
    ("• Dinner Revenue is set to 0 (outsourced, no Aura share). Editable if terms change.", BODY_FONT),
    ("• All currency values are in CHF.", BODY_FONT),
]

for i, (text, font) in enumerate(content, start=1):
    cell = ins.cell(row=i, column=1)
    cell.value = text
    if font:
        cell.font = font

ins.protection.sheet = True
ins.protection.password = "2027"
ins.protection.enable()

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# SHEET PROTECTION (password: 2027)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
PASSWORD = "2027"

for sheet_name in ["Dashboard", "P&L", "Revenue", "Cost"]:
    sheet = wb[sheet_name]
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column):
        for cell in row:
            cell.protection = Protection(locked=True)
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column):
        for cell in row:
            if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
                rgb = str(cell.fill.start_color.rgb)
                if rgb in ("FFFFFF00", "00FFFF00"):
                    cell.protection = Protection(locked=False)
    sheet.protection.sheet = True
    sheet.protection.password = PASSWORD
    sheet.protection.enable()
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = False

wb.save(DST)
print("Done! Saved to", DST)
print(f"New row mappings: FB Total={FB_TOTAL_ROW}, Dinner={DINNER_ROW}, Other={OTHER_ROW}, Annual={ANNUAL_ROW}")
